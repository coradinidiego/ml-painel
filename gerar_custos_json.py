"""
Regenera fixo_por_sku.json e custo_por_sku.json a partir das planilhas de precificação.
Chamado automaticamente pelo endpoint POST /apuracao-bling/recarregar-custos
"""
import json, sys, subprocess
try:
    import openpyxl
except ImportError:
    # Tenta instalar openpyxl para o python atual
    for pip_args in (
        [sys.executable, '-m', 'pip', 'install', 'openpyxl', '--user', '-q'],
        [sys.executable, '-m', 'pip', 'install', 'openpyxl', '-q'],
        ['pip3', 'install', 'openpyxl', '--user', '-q'],
    ):
        try:
            subprocess.check_call(pip_args, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            break
        except Exception:
            continue
    import openpyxl
from pathlib import Path

import os
DIR          = Path(os.environ.get('DATA_DIR') or Path(__file__).parent)
PLANILHA_Z   = DIR / 'Precificacao Comercial Zagodini.xlsx'
PLANILHA_C   = DIR / 'Precificacao Atelie Com S2 de Mãe.xlsx'
FIXO_SAIDA   = DIR / 'fixo_por_sku.json'
CUSTO_SAIDA  = DIR / 'custo_por_sku.json'


def ler_fixo_aba(wb, aba_nome, tem_frete_gratis=True):
    """Col A=SKU, G=Custo Fixo, H=Frete Gratis"""
    import datetime as _dt
    if aba_nome not in wb.sheetnames:
        return {}
    ws = wb[aba_nome]
    resultado = {}
    for row in ws.iter_rows(values_only=True):
        sku = row[0]
        cf  = row[6] if len(row) > 6 else None
        fg  = row[7] if (tem_frete_gratis and len(row) > 7) else None
        if sku is None or cf is None:
            continue
        if isinstance(sku, (_dt.datetime, _dt.date)):
            sku = f'{sku.month}-{sku.year}'
            print(f'AVISO [{aba_nome}]: SKU lido como data pelo Excel, interpretado como "{sku}"', file=sys.stderr)
        sku_s = str(sku).strip()
        if sku_s in ('Cod', 'cod', 'COD', '') or not sku_s:
            continue
        try:
            cf_f = float(cf)
        except (ValueError, TypeError):
            continue
        try:
            fg_f = float(fg) if fg is not None else None
        except (ValueError, TypeError):
            fg_f = None
        if tem_frete_gratis:
            resultado[sku_s] = {'fixo': cf_f, 'freteGratis': fg_f}
        else:
            resultado[sku_s] = cf_f
    return resultado


def ler_custo_aba(wb, aba_nome):
    """Col A=SKU, B=Qtd, C=Produto, D=Custo Un, E=Custo total — usa E (custo total do SKU como vendido).
    SKUs duplicados com custos diferentes: mantém o primeiro e imprime aviso (não trava)."""
    import datetime as _dt
    if aba_nome not in wb.sheetnames:
        return {}
    ws = wb[aba_nome]
    resultado = {}  # SKU → custo total do SKU
    avisos = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        sku      = row[0]
        custo_un = row[4] if len(row) > 4 else None  # col E = Custo total
        if sku is None:
            continue
        # Excel às vezes interpreta SKUs como datas (ex: "1-2755" → datetime(2755,1,1))
        if isinstance(sku, (_dt.datetime, _dt.date)):
            sku = f'{sku.month}-{sku.year}'
            print(f'AVISO [{aba_nome}] Linha {i}: SKU lido como data, interpretado como "{sku}"', file=sys.stderr)
        sku_s = str(sku).strip()
        if sku_s in ('Cod', 'cod', 'COD', '') or not sku_s:
            continue
        try:
            cu = round(float(custo_un), 4)
        except (ValueError, TypeError):
            continue
        if sku_s in resultado:
            if resultado[sku_s] != cu:
                avisos.append(
                    f'AVISO [{aba_nome}] Linha {i}: SKU "{sku_s}" duplicado '
                    f'(R${resultado[sku_s]} vs R${cu}) — mantendo R${resultado[sku_s]}'
                )
        else:
            resultado[sku_s] = cu
    for a in avisos:
        print(a, file=sys.stderr)
    return resultado


def carregar(planilha):
    try:
        return openpyxl.load_workbook(planilha, read_only=True, data_only=True)
    except Exception as e:
        print(f'ERRO ao abrir {planilha}: {e}', file=sys.stderr)
        return None


wb_z = carregar(PLANILHA_Z)
wb_c = carregar(PLANILHA_C)

# fixo_por_sku.json
fixo = {
    'zagodini': {
        'amazon': ler_fixo_aba(wb_z, 'Amazon', tem_frete_gratis=True)  if wb_z else {},
        'shein':  ler_fixo_aba(wb_z, 'Shein',  tem_frete_gratis=False) if wb_z else {},
    },
    'camae': {
        'amazon': ler_fixo_aba(wb_c, 'Amazon', tem_frete_gratis=True)  if wb_c else {},
        'shein':  ler_fixo_aba(wb_c, 'Shein',  tem_frete_gratis=False) if wb_c else {},
    },
}
FIXO_SAIDA.write_text(json.dumps(fixo, ensure_ascii=False, indent=2))
print(f'fixo_por_sku.json: Z-amz={len(fixo["zagodini"]["amazon"])} Z-sh={len(fixo["zagodini"]["shein"])} C-amz={len(fixo["camae"]["amazon"])} C-sh={len(fixo["camae"]["shein"])}')

# custo_por_sku.json
custo = {
    'zagodini': {
        'amazon': ler_custo_aba(wb_z, 'Amazon') if wb_z else {},
        'shein':  ler_custo_aba(wb_z, 'Shein')  if wb_z else {},
    },
    'camae': {
        'amazon': ler_custo_aba(wb_c, 'Amazon') if wb_c else {},
        'shein':  ler_custo_aba(wb_c, 'Shein')  if wb_c else {},
    },
}
CUSTO_SAIDA.write_text(json.dumps(custo, ensure_ascii=False, indent=2))
print(f'custo_por_sku.json: Z-amz={len(custo["zagodini"]["amazon"])} Z-sh={len(custo["zagodini"]["shein"])} C-amz={len(custo["camae"]["amazon"])} C-sh={len(custo["camae"]["shein"])}')
